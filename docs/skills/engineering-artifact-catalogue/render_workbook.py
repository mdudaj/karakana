"""Render this research package's reviewed JSON snapshot to a new Excel file.

This is a documentation helper, not the proposed engineering-workbook tool.
Requires openpyxl in the selected Python environment. Existing files are never
overwritten; reconcile human edits before rendering a subsequent snapshot.
"""

from __future__ import annotations

import argparse
import json
import math
import textwrap
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = Path(__file__).resolve().parent
REVIEW_VALUES = [
    "Not reviewed",
    "Agree with proposal",
    "Requires revision",
    "Needs decision",
]


def render(source: Path, output: Path) -> None:
    if output.exists():
        raise FileExistsError(
            f"Refusing to overwrite {output}. Reconcile review edits and use a new filename."
        )
    data = json.loads(source.read_text(encoding="utf-8"))
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Engineering artifact catalogue: research and update plan"
    workbook.properties.subject = "Draft proposal; no catalogue implementation or approval implied"
    workbook.properties.creator = "Karakana documentation research"
    workbook.properties.description = "Evidence register and proposed global engineering documentation update"
    for number, spec in enumerate(data["sheets"]):
        headers = spec["headers"]
        if not headers or not all(isinstance(h, str) and h for h in headers):
            raise ValueError(f"Invalid headers in {spec['name']}")
        if any(len(row) != len(headers) for row in spec["rows"]):
            raise ValueError(f"Row length mismatch in {spec['name']}")
        sheet = workbook.create_sheet(spec["name"])
        sheet["A1"] = spec["name"]
        sheet["A1"].font = Font(name="Arial", bold=True, size=15, color="16324F")
        sheet["A1"].alignment = Alignment(vertical="center", wrap_text=True)
        sheet["B1"] = "Draft for review"
        sheet["B1"].font = Font(name="Arial", italic=True, color="495766", size=11)
        sheet["A2"] = "Purpose"
        sheet["B2"] = spec["purpose"]
        sheet["A3"] = "Navigation"
        sheet["B3"] = "Return to Read Me"
        sheet["B3"].hyperlink = "#'00 Read Me'!A1"
        sheet["B3"].font = Font(name="Arial", color="155CA8", underline="single", size=11)
        for row in (2, 3):
            for cell in sheet[row]:
                if cell.coordinate != "B3":
                    cell.font = Font(name="Arial", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        widths = [min(65, max(13, w)) for w in (spec.get("widths") or [30] * len(headers))]
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.row_dimensions[1].height = 58
        sheet.row_dimensions[2].height = max(32, 15 * math.ceil(len(spec["purpose"]) / max(1, widths[1] - 2)))
        sheet.row_dimensions[3].height = 21
        for column, header in enumerate(headers, 1):
            cell = sheet.cell(4, column, header)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", fgColor="16324F")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[4].height = 34
        for row_number, row in enumerate(spec["rows"], 5):
            line_count = 1
            for column, value in enumerate(row, 1):
                cell = sheet.cell(row_number, column, value)
                if isinstance(value, str):
                    # Records are literal text, including leading formula-like characters.
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=11, color="172B3A")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if headers[column - 1] in {"Review Decision", "Reviewer Comments"}:
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
                if isinstance(value, str) and value.startswith("https://"):
                    cell.hyperlink = value
                    cell.hyperlink.tooltip = f"Open source for {row[0]}"
                    cell.font = Font(name="Arial", size=11, color="155CA8", underline="single")
                if isinstance(value, str):
                    lines = sum(max(1, len(textwrap.wrap(part, width=max(10, int(widths[column - 1]) - 3))))
                                for part in value.splitlines() or [""])
                    line_count = max(line_count, lines)
            sheet.row_dimensions[row_number].height = min(390, max(32, line_count * 15 + 9))
        final_row = 4 + len(spec["rows"])
        table = Table(displayName=f"EngineeringResearch{number:02d}",
                      ref=f"A4:{get_column_letter(len(headers))}{final_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
        if "Review Decision" in headers:
            column = get_column_letter(headers.index("Review Decision") + 1)
            validation = DataValidation(type="list", formula1='"' + ",".join(REVIEW_VALUES) + '"',
                                        allow_blank=False)
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            validation.errorTitle = "Use a review decision"
            validation.error = "Choose one of the review decisions in the Read Me guide."
            validation.showInputMessage = True
            validation.promptTitle = "Proposal review"
            validation.prompt = "A review response is not deployment permission or test evidence."
            sheet.add_data_validation(validation)
            validation.add(f"{column}5:{column}{final_row}")
        sheet.freeze_panes = "B5"
        sheet.sheet_view.zoomScale = 85
        sheet.print_title_rows = "1:4"
        sheet.print_title_cols = "A:A"
        sheet.print_area = f"A1:{get_column_letter(len(headers))}{final_row}"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        # Wide research registers need horizontal pages to retain readable type.
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 2 if len(headers) > 5 else 1
        sheet.page_setup.fitToHeight = 0
        sheet.oddFooter.center.text = "Draft research and proposed update | Page &P of &N"
        sheet.oddFooter.center.size = 9
    workbook.save(output)
    print(f"Created {output}: {len(workbook.sheetnames)} sheets")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=BASE / "research-and-update-plan.json")
    parser.add_argument("--output", type=Path, default=BASE / "research-and-update-plan.xlsx")
    args = parser.parse_args()
    render(args.source, args.output)
