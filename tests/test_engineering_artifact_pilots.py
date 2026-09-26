"""Observed share-pack outcomes for the controlled synthetic authoring pilot."""

from pathlib import Path
import hashlib
import json
import shutil

import pytest
from openpyxl import load_workbook
from openpyxl.comments import Comment

from karakana.tools.engineering_artifacts import load_bundle, export_workbook, feedback_report


PILOT = Path(__file__).resolve().parents[1] / "docs/skills/engineering-artifact-catalogue/pilots"
PROFILES = ("business", "engineering", "qa", "operations", "executive")
SOURCES = [Path(name) for name in ("requirements.md", "design.md", "plan.md", "release.md")]


@pytest.fixture
def pilot(tmp_path):
    # Include the machine/intent files so missing external context is observable.
    for path in PILOT.iterdir():
        if path.is_file():
            shutil.copy2(path, tmp_path / path.name)
    return tmp_path


def cell(workbook, table, identity, field):
    sheet = workbook[table]
    headers = [item.value for item in sheet[4]]
    row = next(i for i in range(5, sheet.max_row + 1) if sheet.cell(i, 3).value == identity)
    return sheet.cell(row, headers.index(field) + 1)


def test_pilot_preserves_cross_file_scope_machine_authority_and_unobserved_tests(pilot):
    bundle = load_bundle(pilot, SOURCES, "pilot-checklist")
    assert not bundle.issues
    assert len(bundle.documents) == 4
    assert all(doc["metadata"]["status"] == "draft" for doc in bundle.documents)
    assert {source["path"] for source in bundle.sources.values()} == {
        *(str(p) for p in SOURCES), "catalogue.schema.json", "catalogue.json", "brief.md"}
    schema = next(s for s in bundle.sources.values() if s["path"] == "catalogue.schema.json")
    assert schema["authority"] == "Machine source" and schema["revision"] == "candidate-0.1"
    assert schema["sha256"] == hashlib.sha256((pilot / schema["path"]).read_bytes()).hexdigest()
    cases = {r["values"]["Criterion ID"] for r in bundle.records if r["table"] == "Test Cases"}
    criteria = {r["record_id"] for r in bundle.records if r["table"] == "Acceptance"}
    assert cases == criteria == {"AC-01", "AC-02", "AC-03"}
    assert {r["values"]["Result"] for r in bundle.records if r["table"] == "Test Runs"} == {"Not run", "Blocked"}
    assert not any(r["table"] == "Release Notes" for r in bundle.records)
    assert all(r["values"]["Status"] == "Proposed" for r in bundle.records if r["table"] == "ADRs")


@pytest.mark.parametrize("profile", PROFILES)
def test_actual_profile_retains_blockers_and_useful_scoped_detail(pilot, profile):
    bundle = load_bundle(pilot, SOURCES, "pilot-checklist")
    raw_before = {p: p.read_bytes() for p in pilot.iterdir() if p.is_file()}
    output, snapshot = pilot / "view.xlsx", pilot / "view.json"
    export_workbook(bundle, profile, output, snapshot, "controlled-test-unaccepted")
    payload = json.loads(snapshot.read_text())["payload"]
    tables = payload["tables"]
    limits = json.dumps(tables["Readiness and Blockers"])
    assert "RISK-01" in limits and "Blocked" in limits and "Not run" in limits
    assert "support/recovery unconfirmed" in limits and "not approval" in limits
    assert payload["profile_version"] == "0.1" and "Release Notes" not in tables
    required = {
        "business": ("Requirements", "Stories", "Acceptance", "Roadmap", "Guidance and Handover"),
        "engineering": ("ADRs", "Design Views", "UX and Contracts", "Implementation Tasks", "Narrative"),
        "qa": ("Acceptance", "Test Cases", "Test Runs", "Release Plan"),
        "operations": ("Release Plan", "Guidance and Handover", "UX and Contracts"),
        "executive": ("Roadmap", "Milestones", "Release Plan"),
    }
    assert set(required[profile]) <= tables.keys()
    if profile == "engineering":
        assert "memory use grows" in json.dumps(tables["Narrative"])
        assert "Proposed" in json.dumps(tables["ADRs"])
    if profile == "executive":
        assert "forecast only" in json.dumps(tables["Roadmap"])
        assert "Implementation Tasks" in payload["omissions"]
    if profile != "engineering":
        assert "Full narrative: use complete Markdown source links" in payload["omissions"]
    assert all(p.read_bytes() == raw for p, raw in raw_before.items())


@pytest.mark.parametrize("stem", [*(f"{p}-v0.1" for p in PROFILES), "business-v0.2"])
def test_committed_audience_pair_is_complete_current_and_immutable(stem):
    output = PILOT / "exports" / f"{stem}.xlsx"
    snapshot = output.with_suffix(".json")
    before = {p: p.read_bytes() for p in (output, snapshot)}
    ledger = json.loads(snapshot.read_text())
    observation = json.loads((PILOT / "qa-observations.json").read_text())["exports"][stem]
    assert observation["workbook_sha256"] == hashlib.sha256(before[output]).hexdigest()
    assert observation["snapshot_sha256"] == hashlib.sha256(before[snapshot]).hexdigest()
    assert observation["payload_sha256"] == ledger["payload_sha256"]
    assert observation["source_files"] == len(ledger["payload"]["sources"]) == 7
    workbook = load_workbook(output)
    assert set(ledger["payload"]["tables"]) == set(workbook.sheetnames)
    assert workbook.sheetnames[:4] == ["Read Me", "Control", "Source Manifest", "Readiness and Blockers"]
    for name, table in ledger["payload"]["tables"].items():
        sheet = workbook[name]
        assert [c.value for c in sheet[4]] == table["headers"]
        assert [[c.value or "" for c in row] for row in sheet.iter_rows(min_row=5)] == table["rows"]
        assert len(sheet.tables) == 1 and sheet.freeze_panes == "D5"
        assert all(not c.comment and c.data_type != "f" for row in sheet for c in row)
        if "Source Document" in table["headers"]:
            column = table["headers"].index("Source Document") + 1
            for i, row in enumerate(table["rows"], 5):
                c = sheet.cell(i, column)
                if c.value:
                    assert c.hyperlink and (output.parent / c.hyperlink.target).resolve().is_file()
    # Same immutable workbook as a returned copy proves the whole baseline and
    # current source comparison, including all declared source hashes.
    report = feedback_report(PILOT, output, snapshot, [output])
    assert not report["changes"] and not report["stale_sources"] and not report["current_source_issues"]
    assert all(p.read_bytes() == raw for p, raw in before.items())


def test_reviewed_amendment_retains_ids_history_and_no_approval(pilot):
    old = load_bundle(pilot, SOURCES, "pilot-checklist")
    new = load_bundle(pilot, [Path("requirements-v0.2.md"), *SOURCES[1:]], "pilot-checklist")
    old_ids = {(r["table"], r["record_id"]) for r in old.records}
    new_ids = {(r["table"], r["record_id"]) for r in new.records}
    assert new_ids - old_ids == {("Change History", "CHANGE-01")}
    assert not old_ids - new_ids
    updated = next(r for r in new.records if r["table"] == "Acceptance" and r["record_id"] == "AC-01")
    assert "including whitespace" in updated["values"]["Measure / Threshold"]
    assert updated["revision"] == "0.2"
    assert all(doc["metadata"]["status"] == "draft" for doc in new.documents)


def test_stale_feedback_and_conflicting_values_require_review_without_source_writes(pilot):
    bundle = load_bundle(pilot, SOURCES, "pilot-checklist")
    output, snapshot = pilot / "view.xlsx", pilot / "view.json"
    export_workbook(bundle, "business", output, snapshot, "pilot-0.1-unaccepted")
    workbook = load_workbook(output)
    cell(workbook, "Acceptance", "AC-01", "Feedback").value = "Clarify whitespace"
    cell(workbook, "Acceptance", "AC-01", "Expected Outcome").comment = Comment("Proposed question", "synthetic label")
    returned = pilot / "returned.xlsx"
    workbook.save(returned)
    source = pilot / "requirements.md"
    source.write_bytes((pilot / "requirements-v0.2.md").read_bytes())
    value_book = load_workbook(output)
    cell(value_book, "Acceptance", "AC-01", "Measure / Threshold").value = "Normalize whitespace"
    contradictory = pilot / "contradictory.xlsx"
    value_book.save(contradictory)
    immutable = {p: p.read_bytes() for p in (source, output, snapshot, returned, contradictory)}
    report = feedback_report(pilot, output, snapshot, [returned, contradictory])
    assert report["stale_sources"] == ["requirements.md"] and not report["current_source_issues"]
    assert any(c.get("status") == "proposed_stale" for c in report["changes"])
    assert any(c.get("status") == "conflict" and c.get("field") == "Measure / Threshold" for c in report["changes"])
    assert any(c["kind"] == "cell_comment" for c in report["changes"])
    assert all(p.read_bytes() == raw for p, raw in immutable.items())
