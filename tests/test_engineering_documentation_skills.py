"""Integration checks for reusable guidance and its authored example sources."""

from pathlib import Path
import hashlib
import json
import re
import shutil

import pytest

from karakana.evals.loader import EvalLoader
from karakana.skills.loader import SkillLoader
from karakana.skills.validator import SkillValidator
from karakana.tools.engineering_artifacts import load_bundle, export_workbook, feedback_report


ROOT = Path(__file__).resolve().parents[1]
SKILLS = (
    "engineering-requirements", "engineering-design-records",
    "engineering-delivery-planning", "engineering-release-documentation",
    "engineering-workbooks",
)


def copy_examples(root):
    paths = []
    for name in SKILLS:
        for source in (ROOT / "skills" / name).glob("examples/*.md"):
            target = root / source.name
            shutil.copyfile(source, target)
            paths.append(target)
    return paths


@pytest.mark.parametrize("name", SKILLS)
def test_catalogue_skills_load_validate_and_have_discoverable_evals(name):
    loader = SkillLoader(ROOT / "skills")
    assert name in loader.list_skills()
    skill = loader.load_skill(name)
    result = SkillValidator().validate(skill.path)
    assert result.is_valid and not result.warnings
    assert skill.status == "experimental" and skill.scope == "bundled"
    cases = EvalLoader(ROOT).load_cases(skill=name)
    assert cases and all(case.input.skill == name for case in cases)
    assert all((ROOT / case.input.prompt_file).is_file() for case in cases)


@pytest.mark.parametrize("name", SKILLS)
def test_conditional_skill_resources_resolve_inside_repository(name):
    for source in (ROOT / "skills" / name).rglob("*.md"):
        for target in re.findall(r"\[[^\]]+\]\(([^)]+)\)", source.read_text()):
            if "://" in target or target.startswith("#"):
                continue
            path = (source.parent / target.split("#", 1)[0]).resolve()
            assert path.is_relative_to(ROOT) and path.is_file(), (source, target)


def test_authored_examples_are_valid_alone_and_combined(tmp_path):
    sources = copy_examples(tmp_path)
    before = {p: p.read_bytes() for p in sources}
    for source in sources:
        assert load_bundle(tmp_path, [source], "example").records
    bundle = load_bundle(tmp_path, sources, "example")
    assert bundle.records and not bundle.issues
    assert all(doc["metadata"]["status"] == "draft" for doc in bundle.documents)
    assert not any(record["table"] == "Release Notes" for record in bundle.records)
    assert all(p.read_bytes() == value for p, value in before.items())


@pytest.mark.parametrize("audience", ["business", "engineering", "qa", "operations", "executive"])
def test_authored_pack_profiles_preserve_uncertainty_and_feedback(tmp_path, audience):
    from openpyxl import load_workbook

    sources = copy_examples(tmp_path)
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}
    output, snapshot = tmp_path / "view.xlsx", tmp_path / "view.json"
    bundle = load_bundle(tmp_path, sources, "example")
    export_workbook(bundle, audience, output, snapshot, "synthetic-authored-0.1")
    payload = json.loads(snapshot.read_text())["payload"]
    limitations = json.dumps(payload["tables"]["Readiness and Blockers"])
    assert "Blocked" in limitations and "RISK-L01" in limitations
    assert "not approval" in limitations
    assert not any(r["values"].get("Status") == "Accepted" for r in bundle.records)
    assert "Release Notes" not in payload["tables"]
    if audience == "engineering":
        assert "Replacing every native store" in json.dumps(payload["tables"]["Narrative"])
        assert "Proposed" in json.dumps(payload["tables"]["ADRs"])
    if audience == "executive":
        assert "forecast only" in json.dumps(payload["tables"]["Roadmap"])
        assert "Implementation Tasks" in payload["omissions"]
    workbook = load_workbook(output)
    ws = workbook["Readiness and Blockers"]
    headers = [cell.value for cell in ws[4]]
    ws.cell(5, headers.index("Feedback") + 1, "Resolve scope before relying on this view.")
    returned = tmp_path / "annotated.xlsx"
    workbook.save(returned)
    immutable = {p: p.read_bytes() for p in (output, snapshot, returned)}
    report = feedback_report(tmp_path, output, snapshot, [returned])
    assert len(report["changes"]) == 1
    assert report["changes"][0]["field"] == "Feedback"
    assert report["changes"][0]["status"] == "proposed"
    assert all(p.read_bytes() == value for p, value in immutable.items())
    assert all(hashlib.sha256(p.read_bytes()).hexdigest() == digest for p, digest in before.items())
