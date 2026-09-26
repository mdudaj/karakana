from pathlib import Path
import json
import shutil

import pytest

from karakana.tools.engineering_artifacts import (
    ArtifactError, load_bundle, export_workbook, feedback_report, propose_bindings,
)


def document(root, statement="The shared view retains its source identity."):
    path = root / "source.md"
    path.write_text(f'''---
contract_version: "0.1"
namespace: "example"
document_id: "DOC-001"
document_type: "document"
content_version: "0.1"
status: "draft"
owner: "unconfirmed"
requested_action: "review"
source_authority: "authored Markdown"
---

# Source

<a id="requirements"></a>
## Requirements

| Requirement ID | Class | Statement | Rationale | Source ID | Proposed Priority | Agreed Priority | Decision Evidence | Owner | Status |
| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |
| REQ-001 | Stakeholder | {statement} | Source context remains available. |  | unconfirmed | unconfirmed |  | unconfirmed | draft |
''')
    return path


def fixture(root):
    source = Path(__file__).resolve().parents[1] / "docs/skills/engineering-artifact-catalogue/examples"
    shutil.copytree(source, root / "examples")
    return root / "examples/feature.md"


def exported(root, sources=None, audience="business"):
    pytest.importorskip("openpyxl")
    sources = sources or [document(root)]
    bundle = load_bundle(root, sources, "example")
    output, snapshot = root / "export.xlsx", root / "export.json"
    export_workbook(bundle, audience, output, snapshot, "fixture-baseline")
    return output, snapshot


def test_markdown_validation_does_not_need_workbook_library(tmp_path):
    bundle = load_bundle(tmp_path, [document(tmp_path)], "example")
    assert bundle.records
    assert bundle.records[0]["record_id"] == "REQ-001"


def test_native_identity_and_source_bytes_preserved(tmp_path):
    source = fixture(tmp_path)
    before = {p: p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    output, snapshot = exported(tmp_path, [source], "engineering")
    data = json.loads(snapshot.read_text())["payload"]
    text = json.dumps(data)
    assert "req-example-story-1" in text and "req-example-issue-1" in text
    assert all(p.read_bytes() == content for p, content in before.items())
    assert len(data["sources"]) == 4
    assert output.is_file()


def test_generic_literal_export_and_no_overwrite(tmp_path):
    from openpyxl import load_workbook
    output, snapshot = exported(tmp_path, [document(tmp_path, "=1+1")])
    ws = load_workbook(output)["Requirements"]
    assert any(c.value == "=1+1" and c.data_type == "s" for row in ws for c in row)
    before = output.read_bytes(), snapshot.read_bytes()
    bundle = load_bundle(tmp_path, [tmp_path / "source.md"], "example")
    with pytest.raises(FileExistsError):
        export_workbook(bundle, "business", output, snapshot, "fixture-baseline")
    assert before == (output.read_bytes(), snapshot.read_bytes())


def test_feedback_three_way_conflict_preserves_inputs(tmp_path):
    from openpyxl import load_workbook
    output, snapshot = exported(tmp_path)
    returned = tmp_path / "review.xlsx"
    wb = load_workbook(output)
    ws = wb["Requirements"]
    headers = [cell.value for cell in ws[4]]
    ws.cell(5, headers.index("Statement") + 1, "Reviewer proposes a different value.")
    wb.save(returned)
    document(tmp_path, "Current source has independently changed.")
    before = {p: p.read_bytes() for p in (output, snapshot, returned, tmp_path / "source.md")}
    report = feedback_report(tmp_path, output, snapshot, [returned])
    assert report["stale_sources"]
    assert any(change["status"] == "conflict" for change in report["changes"])
    assert all(p.read_bytes() == content for p, content in before.items())


@pytest.mark.parametrize("mutation", ["namespace", "duplicate_key", "version", "duplicate_record"])
def test_invalid_sources_reject(tmp_path, mutation):
    path = document(tmp_path)
    text = path.read_text()
    if mutation == "namespace":
        text = text.replace('namespace: "example"', 'namespace: ""')
    elif mutation == "duplicate_key":
        text = text.replace('namespace: "example"', 'namespace: "example"\nnamespace: "example"')
    elif mutation == "version":
        text = text.replace('contract_version: "0.1"', 'contract_version: "999"')
    else:
        text += next(line for line in text.splitlines() if line.startswith("| REQ-001 |")) + "\n"
    path.write_text(text)
    with pytest.raises(ArtifactError):
        load_bundle(tmp_path, [path], "example")


@pytest.mark.parametrize("audience", ["business", "engineering", "qa", "operations", "executive"])
def test_audience_blockers_and_omissions(tmp_path, audience):
    source = fixture(tmp_path)
    release = tmp_path / "examples/release.md"
    output, snapshot = exported(tmp_path, [source, release], audience)
    from openpyxl import load_workbook
    wb = load_workbook(output)
    values = json.dumps([[c.value for c in row] for row in wb["Readiness and Blockers"]])
    assert "RISK-001" in values and "Not run" in values and "unverified" in values
    payload = json.loads(snapshot.read_text())["payload"]
    assert payload["audience"] == audience and payload["sources"]
    if audience != "engineering":
        assert "Source Bindings" not in wb.sheetnames
        assert payload["omissions"]
    assert wb["Control"].freeze_panes == "D5"


@pytest.mark.parametrize("mutation", ["text", "metadata", "duplicate", "namespace"])
def test_engine_source_drift_and_ambiguity_block(tmp_path, mutation):
    source = fixture(tmp_path)
    engine = tmp_path / "examples/engine-fixture/prd.json"
    data = json.loads(engine.read_text())
    if mutation == "text":
        data["functional_requirements"][0] = "Changed native statement"
    elif mutation == "metadata":
        data["metadata"]["changed"] = True
    elif mutation == "duplicate":
        data["functional_requirements"] *= 2
    else:
        data["project"] = "another-namespace"
    engine.write_text(json.dumps(data))
    with pytest.raises(ArtifactError):
        load_bundle(tmp_path, [source], "example")


def test_binding_proposal_preserves_ids_but_never_resolves_changes(tmp_path):
    source = fixture(tmp_path)
    engine = tmp_path / "examples/engine-fixture/prd.json"
    before = engine.read_bytes(), source.read_bytes()
    proposal = propose_bindings(tmp_path, engine, "example", registry=source,
                                output=tmp_path / "proposal.md")
    assert "REQ-001" in proposal and "AC-001" in proposal
    assert "| resolved |" not in proposal and "| proposed |" in proposal
    assert before == (engine.read_bytes(), source.read_bytes())
    data = json.loads(engine.read_text())
    data["functional_requirements"][0] = "Edited criterion intent needs association review"
    engine.write_text(json.dumps(data))
    edited = propose_bindings(tmp_path, engine, "example", registry=source)
    assert "REQ-001" in edited and "| unresolved |" in edited


def test_duplicate_bootstrap_proposals_are_unresolved(tmp_path):
    source = fixture(tmp_path)
    engine = tmp_path / "examples/engine-fixture/prd.json"
    data = json.loads(engine.read_text())
    data["functional_requirements"] *= 2
    engine.write_text(json.dumps(data))
    result = propose_bindings(tmp_path, engine, "example")
    assert result.count("| unresolved |") == 2


def test_removed_binding_retained_for_review(tmp_path):
    source = fixture(tmp_path)
    engine = tmp_path / "examples/engine-fixture/prd.json"
    data = json.loads(engine.read_text())
    data["functional_requirements"] = []
    engine.write_text(json.dumps(data))
    result = propose_bindings(tmp_path, engine, "example", registry=source)
    assert "REQ-001" in result and "| unresolved |" in result


@pytest.mark.parametrize("same", [True, False])
def test_output_alias_and_escape_refused_without_writes(tmp_path, same):
    bundle = load_bundle(tmp_path, [document(tmp_path)], "example")
    output = tmp_path / "same.xlsx"
    snapshot = output if same else tmp_path.parent / "escape-output.json"
    with pytest.raises(ArtifactError):
        export_workbook(bundle, "business", output, snapshot, "fixture-baseline")
    assert not output.exists() and not snapshot.exists()


def test_existing_snapshot_prevents_partial_workbook(tmp_path):
    bundle = load_bundle(tmp_path, [document(tmp_path)], "example")
    snapshot = tmp_path / "existing.json"
    snapshot.write_text("review evidence")
    with pytest.raises(FileExistsError):
        export_workbook(bundle, "business", tmp_path / "new.xlsx", snapshot, "fixture-baseline")
    assert not (tmp_path / "new.xlsx").exists() and snapshot.read_text() == "review evidence"


@pytest.mark.parametrize("mutation", ["snapshot", "baseline"])
def test_tampered_baseline_rejected(tmp_path, mutation):
    from openpyxl import load_workbook
    output, snapshot = exported(tmp_path)
    returned = tmp_path / "review.xlsx"
    shutil.copyfile(output, returned)
    if mutation == "snapshot":
        data = json.loads(snapshot.read_text())
        data["payload"]["baseline_revision"] = "forged"
        snapshot.write_text(json.dumps(data))
    else:
        wb = load_workbook(output)
        wb["Read Me"]["D5"] = "changed"
        wb.save(output)
    with pytest.raises(ArtifactError):
        feedback_report(tmp_path, output, snapshot, [returned])


def test_formula_unknown_sheet_comments_and_row_edits_reported(tmp_path):
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    output, snapshot = exported(tmp_path)
    returned = tmp_path / "review.xlsx"
    wb = load_workbook(output)
    ws = wb["Requirements"]
    headers = [c.value for c in ws[4]]
    ws.cell(5, headers.index("Statement") + 1, "=1+1")
    ws.cell(5, headers.index("Feedback") + 1).comment = Comment("Keep source identity", "fixture reviewer")
    ws.append(["example", "requirement", "NEW-001", *([""] * (len(headers) - 3))])
    wb.create_sheet("Extra reviewer notes")["A1"] = "preserve this unknown content"
    wb.save(returned)
    before = returned.read_bytes()
    report = feedback_report(tmp_path, output, snapshot, [returned])
    kinds = {change["kind"] for change in report["changes"]}
    assert {"formula_input", "worksheet_added", "cell_comment", "row_added"} <= kinds
    assert before == returned.read_bytes()


def test_conflicting_returned_workbooks(tmp_path):
    from openpyxl import load_workbook
    output, snapshot = exported(tmp_path)
    returned = []
    for i in range(2):
        wb = load_workbook(output)
        ws = wb["Requirements"]
        headers = [c.value for c in ws[4]]
        ws.cell(5, headers.index("Statement") + 1, "Proposal " + str(i))
        path = tmp_path / f"review-{i}.xlsx"
        wb.save(path)
        returned.append(path)
    report = feedback_report(tmp_path, output, snapshot, returned)
    assert len(report["changes"]) == 2
    assert all(change["status"] == "conflict" for change in report["changes"])


def test_missing_current_source_is_reported(tmp_path):
    output, snapshot = exported(tmp_path)
    returned = tmp_path / "review.xlsx"
    shutil.copyfile(output, returned)
    (tmp_path / "source.md").unlink()
    report = feedback_report(tmp_path, output, snapshot, [returned])
    assert report["stale_sources"] == ["source.md"]
    assert report["current_source_issues"]


def test_source_escape_and_symlink_escape(tmp_path):
    outside = tmp_path.parent / (tmp_path.name + "-outside.md")
    outside.write_text("not inside the root")
    try:
        with pytest.raises(ArtifactError):
            load_bundle(tmp_path, [outside], "example")
        alias = tmp_path / "alias.md"
        alias.symlink_to(outside)
        with pytest.raises(ArtifactError):
            load_bundle(tmp_path, [alias], "example")
    finally:
        outside.unlink()


def test_unanchored_table_and_false_acceptance_refused(tmp_path):
    path = document(tmp_path)
    path.write_text(path.read_text().replace('<a id="requirements"></a>', ""))
    with pytest.raises(ArtifactError):
        load_bundle(tmp_path, [path], "example")
    path = document(tmp_path)
    path.write_text(path.read_text().replace('status: "draft"', 'status: "accepted"'))
    with pytest.raises(ArtifactError):
        load_bundle(tmp_path, [path], "example")


def test_false_test_pass_refused(tmp_path):
    source = fixture(tmp_path)
    source.write_text(source.read_text().replace("| Not run |", "| Pass |"))
    with pytest.raises(ArtifactError):
        load_bundle(tmp_path, [source], "example")


def test_source_changes_after_validation_refuse_export(tmp_path):
    source = document(tmp_path)
    bundle = load_bundle(tmp_path, [source], "example")
    source.write_text(source.read_text() + "\nChanged after validation\n")
    with pytest.raises(ArtifactError):
        export_workbook(bundle, "business", tmp_path / "out.xlsx", tmp_path / "out.json", "fixture")
    assert not (tmp_path / "out.xlsx").exists()


def test_metadata_secrets_are_not_exported(tmp_path):
    source = fixture(tmp_path)
    engine = tmp_path / "examples/engine-fixture/prd.json"
    data = json.loads(engine.read_text())
    data["metadata"]["api_key"] = "FAKE-SECRET-DO-NOT-EXPORT"
    engine.write_text(json.dumps(data, indent=2, sort_keys=True) + "\n")
    import hashlib
    digest = hashlib.sha256(engine.read_bytes()).hexdigest()
    text = source.read_text()
    old = next(line for line in text.splitlines() if line.startswith("| req-example |"))
    old_hash = old.split("|")[9].strip()
    source.write_text(text.replace(old_hash, digest))
    _, snapshot = exported(tmp_path, [source], "engineering")
    assert "FAKE-SECRET-DO-NOT-EXPORT" not in snapshot.read_text()


def test_lazy_spreadsheet_dependency(tmp_path, monkeypatch):
    import builtins
    original = builtins.__import__
    def guarded(name, *args, **kwargs):
        if name == "openpyxl" or name.startswith("openpyxl."):
            raise ImportError("intentionally unavailable")
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, "__import__", guarded)
    bundle = load_bundle(tmp_path, [document(tmp_path)], "example")
    assert bundle.records
    with pytest.raises(ArtifactError, match="optional workbooks"):
        export_workbook(bundle, "business", tmp_path / "out.xlsx", tmp_path / "out.json", "fixture")
    assert not (tmp_path / "out.xlsx").exists()


def test_binding_reorder_retains_persisted_identity_as_proposal(tmp_path):
    source = fixture(tmp_path)
    engine = tmp_path / "examples/engine-fixture/prd.json"
    data = json.loads(engine.read_text())
    data["functional_requirements"].insert(0, "Another source record")
    engine.write_text(json.dumps(data))
    proposal = propose_bindings(tmp_path, engine, "example", registry=source)
    line = next(line for line in proposal.splitlines() if "| REQ-001 |" in line)
    assert "| [1] |" in line and "| proposed |" in line


@pytest.mark.parametrize("marker", ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"])
def test_unsupported_transport_preserves_returned_file(tmp_path, marker):
    import zipfile
    output, snapshot = exported(tmp_path)
    returned = tmp_path / "unsupported.xlsx"
    shutil.copyfile(output, returned)
    with zipfile.ZipFile(returned, "a") as archive:
        archive.writestr(marker, "synthetic unsupported feature")
    before = returned.read_bytes()
    with pytest.raises(ArtifactError, match="unsupported"):
        feedback_report(tmp_path, output, snapshot, [returned])
    assert returned.read_bytes() == before


def test_sorted_rows_create_no_feedback(tmp_path):
    from openpyxl import load_workbook
    source = document(tmp_path)
    line = next(line for line in source.read_text().splitlines() if line.startswith("| REQ-001 |"))
    source.write_text(source.read_text() + line.replace("REQ-001", "REQ-002") + "\n")
    output, snapshot = exported(tmp_path, [source])
    wb = load_workbook(output)
    ws = wb["Requirements"]
    first, second = [c.value for c in ws[5]], [c.value for c in ws[6]]
    for i in range(len(first)):
        ws.cell(5, i + 1, second[i]); ws.cell(6, i + 1, first[i])
    returned = tmp_path / "sorted.xlsx"; wb.save(returned)
    assert feedback_report(tmp_path, output, snapshot, [returned])["changes"] == []


def test_removal_and_modification_conflict_across_workbooks(tmp_path):
    from openpyxl import load_workbook
    output, snapshot = exported(tmp_path)
    deletion = load_workbook(output); deletion["Requirements"].delete_rows(5)
    deleted = tmp_path / "deleted.xlsx"; deletion.save(deleted)
    modification = load_workbook(output); ws = modification["Requirements"]
    headers = [c.value for c in ws[4]]; ws.cell(5, headers.index("Statement") + 1, "Changed proposal")
    modified = tmp_path / "modified.xlsx"; modification.save(modified)
    report = feedback_report(tmp_path, output, snapshot, [deleted, modified])
    assert {change["kind"] for change in report["changes"]} == {"row_removed", "field_change"}
    assert all(change["status"] == "conflict" for change in report["changes"])


def test_cli_exports_and_feedback(tmp_path, capsys):
    from karakana.tools.engineering_artifacts import main
    source = document(tmp_path)
    assert main(["--root", str(tmp_path), "validate", "--source", str(source), "--namespace", "example"]) == 0
    assert main(["--root", str(tmp_path), "export", "--source", str(source), "--namespace", "example",
                 "--audience", "executive", "--output", "cli.xlsx", "--snapshot", "cli.json", "--baseline-revision", "fixture"]) == 0
    shutil.copyfile(tmp_path / "cli.xlsx", tmp_path / "returned.xlsx")
    assert main(["--root", str(tmp_path), "feedback", "--baseline", "cli.xlsx", "--snapshot", "cli.json",
                 "--returned", "returned.xlsx", "--output", "feedback.json"]) == 0
    assert json.loads((tmp_path / "feedback.json").read_text())["changes"] == []
    assert "no changes applied" in capsys.readouterr().out


def test_secret_paths_rejected_before_read(tmp_path):
    path = tmp_path / ".env"
    path.write_text("SYNTHETIC_CREDENTIAL=fixture")
    with pytest.raises(ArtifactError, match="outside this tool"):
        load_bundle(tmp_path, [path], "example")


def test_declared_engine_readiness_is_a_recorded_limit(tmp_path):
    source = document(tmp_path)
    readiness = tmp_path / "readiness.json"
    readiness.write_text(json.dumps({"req_id": "req-example", "status": "blocked", "ready": False}))
    import hashlib
    from karakana.tools.engineering_artifacts import TABLE_HEADERS
    headers = TABLE_HEADERS["Source Manifest"]
    row = ["SRC-READY", "Engine JSON", "readiness.json", "", "fixture", hashlib.sha256(readiness.read_bytes()).hexdigest(), "Engine JSON", "readiness.json"]
    text = '\n<a id="source-manifest"></a>\n## Source Manifest\n\n'
    text += '| ' + ' | '.join(headers) + ' |\n| ' + ' | '.join(['---'] * len(headers)) + ' |\n| ' + ' | '.join(row) + ' |\n'
    source.write_text(source.read_text() + text)
    _, snapshot = exported(tmp_path, [source], "executive")
    payload = json.loads(snapshot.read_text())["payload"]
    assert "status=blocked" in json.dumps(payload["tables"]["Readiness and Blockers"])
    assert "not approval" in json.dumps(payload["tables"]["Readiness and Blockers"])
    assert next(item for item in payload["sources"] if item["path"] == "readiness.json")["revision"] == "fixture"


def test_unknown_manifest_authority_blocks_export(tmp_path):
    source = document(tmp_path)
    from karakana.tools.engineering_artifacts import TABLE_HEADERS
    import hashlib
    machine = tmp_path / "schema.json"; machine.write_text('{}')
    headers = TABLE_HEADERS["Source Manifest"]
    row = ["SRC-001", "Machine source", "schema.json", "", "v1", hashlib.sha256(machine.read_bytes()).hexdigest(), "unconfirmed", "schema.json"]
    source.write_text(source.read_text() + '\n<a id="manifest"></a>\n## Source Manifest\n\n| ' + ' | '.join(headers) + ' |\n| ' + ' | '.join(['---'] * len(headers)) + ' |\n| ' + ' | '.join(row) + ' |\n')
    with pytest.raises(ArtifactError, match="authority"):
        load_bundle(tmp_path, [source], "example")


def test_returned_row_deletion_conflicts_with_current_source_edit(tmp_path):
    from openpyxl import load_workbook
    output, snapshot = exported(tmp_path)
    wb = load_workbook(output); wb["Requirements"].delete_rows(5)
    returned = tmp_path / "deleted.xlsx"; wb.save(returned)
    document(tmp_path, "Current source independently changed this row.")
    report = feedback_report(tmp_path, output, snapshot, [returned])
    assert any(change["kind"] == "row_removed" and change["status"] == "conflict" for change in report["changes"])


def test_returned_added_id_collision_with_current_source_conflicts(tmp_path):
    from openpyxl import load_workbook
    output, snapshot = exported(tmp_path)
    wb = load_workbook(output); ws = wb["Requirements"]
    values = [c.value for c in ws[5]]
    values[2] = values[3] = "REQ-002"
    ws.append(values)
    returned = tmp_path / "added.xlsx"; wb.save(returned)
    source = tmp_path / "source.md"
    row = next(line for line in source.read_text().splitlines() if line.startswith("| REQ-001 |"))
    source.write_text(source.read_text() + row.replace("REQ-001", "REQ-002").replace("The shared view retains its source identity.", "A distinct current source need.") + '\n')
    report = feedback_report(tmp_path, output, snapshot, [returned])
    assert any(change["kind"] == "row_added" and change["status"] == "conflict" for change in report["changes"])


def test_engineering_retains_intro_and_structured_section_context(tmp_path):
    source = document(tmp_path)
    source.write_text(source.read_text().replace('# Source', '# Source\n\nIntroductory uncertainty.') + '\nA qualification after the table.\n')
    _, snapshot = exported(tmp_path, [source], 'engineering')
    narratives = json.loads(snapshot.read_text())['payload']['tables']['Narrative']
    rendered = json.dumps(narratives)
    assert 'Introductory uncertainty.' in rendered
    assert 'A qualification after the table.' in rendered
    assert {row[2] for row in narratives['rows']} == {'DOC-001#', 'DOC-001#requirements'}


def test_oversized_snapshot_is_refused_without_unusable_outputs(tmp_path, monkeypatch):
    from karakana.tools import engineering_artifacts as tool
    source = document(tmp_path); bundle = load_bundle(tmp_path, [source], 'example')
    monkeypatch.setattr(tool, 'MAX_SOURCE_BYTES', 1500)
    with pytest.raises(ArtifactError, match='snapshot size'):
        export_workbook(bundle, 'engineering', tmp_path/'large.xlsx', tmp_path/'large.json', 'fixture')
    assert not (tmp_path/'large.xlsx').exists() and not (tmp_path/'large.json').exists()
