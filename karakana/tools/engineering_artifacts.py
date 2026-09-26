"""Opt-in local Markdown validation, derived XLSX exports and feedback reports.

No source apply operation, engine regeneration, network or model execution.
Run ``python -m karakana.tools.engineering_artifacts --help`` for the module CLI.
Spreadsheet support is lazy-loaded from the optional ``workbooks`` extra.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass, field
from datetime import datetime, timezone
import hashlib
import io
import json
import os
from pathlib import Path
import re
import uuid
import zipfile

import yaml

from karakana.requirements.schemas import RequirementPRD, UserStory, IssueDraft, ReadinessCheck


class ArtifactError(ValueError):
    """Source or transport does not satisfy the supported contract."""


CONTRACT_VERSION = "0.1"
MAX_SOURCE_BYTES = 2_000_000
MAX_WORKBOOK_BYTES = 20_000_000
MAX_WORKBOOK_SCAN_CELLS = 500_000
META_FIELDS = {"contract_version", "namespace", "document_id", "document_type",
               "content_version", "status", "owner", "requested_action", "source_authority"}
DOCUMENT_STATUSES = {"draft", "in_review", "accepted", "superseded", "withdrawn"}
RECORD_TYPES = {"prd", "requirement", "story", "criterion", "adr", "design", "contract",
                "outcome", "milestone", "task", "release", "release_note", "test_case",
                "test_run", "risk", "review", "evidence", "guide", "change", "source",
                "binding", "link", "document"}

# Exact versioned headers are populated below from the reviewed generic contract.
TABLE_HEADERS = {'Artifact Register': ['Artifact ID',
                       'Type',
                       'Title',
                       'Status',
                       'Owner',
                       'Markdown Path',
                       'Source Section',
                       'Source Revision',
                       'Source SHA-256',
                       'Source Authority',
                       'Supersedes ID',
                       'Machine Source Link'],
 'Sources': ['Source ID',
             'Publisher / Owner',
             'URI / Path',
             'Edition / Revision',
             'Inspected Date',
             'Supported Claim',
             'Limitations'],
 'Requirements': ['Requirement ID',
                  'Class',
                  'Statement',
                  'Rationale',
                  'Source ID',
                  'Proposed Priority',
                  'Agreed Priority',
                  'Decision Evidence',
                  'Owner',
                  'Status'],
 'Stories': ['Story ID',
             'Actor',
             'Goal',
             'Benefit',
             'Conversation / Scope',
             'Dependency IDs',
             'Estimate / Basis',
             'Owner',
             'Status'],
 'Acceptance': ['Criterion ID',
                'Context',
                'Event',
                'Expected Outcome',
                'Measure / Threshold',
                'Verification Method',
                'Source Artifact ID',
                'Owner'],
 'ADRs': ['ADR ID',
          'Title',
          'Status',
          'Decision Summary',
          'Rationale Summary',
          'Consequences Summary',
          'Decision Owner',
          'Source Artifact ID',
          'Superseded By',
          'Decision Evidence'],
 'ADR Source Sections': ['ADR ID',
                         'Section',
                         'Sequence',
                         'Summary / Excerpt',
                         'Source Artifact ID',
                         'Source Section',
                         'Source Revision'],
 'Design Views': ['Design ID',
                  'View ID',
                  'Stakeholder / Concern',
                  'Viewpoint / Scope',
                  'Elements / Interfaces',
                  'Scenario',
                  'Quality Constraint',
                  'Diagram / Schema Link',
                  'Source Artifact ID'],
 'UX and Contracts': ['Element ID',
                      'Behavior / States',
                      'Copy / Interaction',
                      'Accessibility',
                      'Shared Components / Tokens',
                      'Contract Link / Version',
                      'Examples / Constraints',
                      'Source Artifact ID'],
 'Roadmap': ['Outcome ID',
             'Audience',
             'Outcome',
             'Success Measure',
             'Horizon',
             'Confidence',
             'Proposal / Commitment',
             'Dependencies',
             'Review Trigger',
             'Source Artifact ID'],
 'Milestones': ['Milestone ID',
                'Goal',
                'Output',
                'Owner',
                'Checkpoint',
                'Acceptance',
                'Estimate / Basis',
                'Forecast / Agreement Evidence',
                'Status',
                'Source Artifact ID'],
 'Implementation Tasks': ['Task ID',
                          'Inspect References / Revision',
                          'Files / Interfaces',
                          'Implementation Steps',
                          'Check / Expected Result',
                          'Dependencies',
                          'Owner',
                          'Approvals / Recovery',
                          'Source Artifact ID'],
 'Release Plan': ['Release ID',
                  'Scope Baseline',
                  'Version Scheme',
                  'Candidate Revision',
                  'Target / Environment',
                  'Readiness Gates',
                  'Rollout',
                  'Recovery',
                  'Support / Owner',
                  'Status'],
 'Release Notes': ['Note ID',
                   'Release ID',
                   'Category',
                   'Audience',
                   'Delivered Change',
                   'Affected Behavior',
                   'Compatibility / Migration',
                   'Known Limits',
                   'Evidence ID',
                   'Source Artifact ID'],
 'Test Cases': ['Test ID',
                'Criterion ID',
                'Preconditions',
                'Steps',
                'Expected Result',
                'Method / Test Source',
                'Owner'],
 'Test Runs': ['Run ID',
               'Test ID',
               'Result',
               'Actual Observation',
               'Revision',
               'Environment',
               'Executed UTC',
               'Executor',
               'Evidence ID'],
 'Risks and Decisions': ['Item ID',
                         'Type',
                         'Uncertainty / Decision Needed',
                         'Impact',
                         'Mitigation / Options',
                         'Trigger / Horizon',
                         'Owner',
                         'Status',
                         'Evidence ID'],
 'Links': ['Link ID',
           'From Namespace',
           'From Type',
           'From ID',
           'Relation',
           'To Namespace',
           'To Type',
           'To ID',
           'Rationale'],
 'Evidence': ['Evidence ID',
              'Related Artifact / Run ID',
              'URI / Path',
              'Revision',
              'Environment',
              'Observed Result',
              'Observed UTC',
              'Limitations',
              'Content Hash'],
 'Reviews and Approvals': ['Review ID',
                           'Source Artifact ID',
                           'Source Revision',
                           'Reviewer / Role',
                           'Review Type',
                           'Review Decision',
                           'Reviewer Comments',
                           'Reviewed UTC',
                           'Evidence ID',
                           'Reconciliation Status'],
 'Change History': ['Change ID',
                    'Artifact ID',
                    'Old Revision',
                    'New Revision',
                    'Reason / Impact',
                    'Author',
                    'Reviewer',
                    'Evidence ID',
                    'Feedback Export ID'],
 'Guidance and Handover': ['Guide ID',
                           'Audience',
                           'Task',
                           'Steps / Guidance Link',
                           'Support / Runbook Link',
                           'Known Limits',
                           'Owner',
                           'Tested Revision'],
 'Source Bindings': ['Binding ID',
                     'Record Type',
                     'Record ID',
                     'Native Owner ID',
                     'Source Path',
                     'Source Kind',
                     'Source Field',
                     'Source Locator',
                     'Source Text SHA-256',
                     'Source File SHA-256',
                     'Binding Status'],
 'Source Manifest': ['Source ID',
                     'Source Kind',
                     'Path',
                     'Selector',
                     'Revision',
                     'SHA-256',
                     'Authority',
                     'Full-Source Link']}
TABLE_IDENTITIES = {
    "Artifact Register": ("document", ("Artifact ID",)),
    "Sources": ("source", ("Source ID",)), "Requirements": ("requirement", ("Requirement ID",)),
    "Stories": ("story", ("Story ID",)), "Acceptance": ("criterion", ("Criterion ID",)),
    "ADRs": ("adr", ("ADR ID",)), "ADR Source Sections": ("adr", ("ADR ID", "Section", "Sequence")),
    "Design Views": ("design", ("Design ID", "View ID")),
    "UX and Contracts": ("contract", ("Element ID",)), "Roadmap": ("outcome", ("Outcome ID",)),
    "Milestones": ("milestone", ("Milestone ID",)),
    "Implementation Tasks": ("task", ("Task ID",)), "Release Plan": ("release", ("Release ID",)),
    "Release Notes": ("release_note", ("Note ID",)), "Test Cases": ("test_case", ("Test ID",)),
    "Test Runs": ("test_run", ("Run ID",)), "Risks and Decisions": ("risk", ("Item ID",)),
    "Links": ("link", ("Link ID",)), "Evidence": ("evidence", ("Evidence ID",)),
    "Reviews and Approvals": ("review", ("Review ID",)), "Change History": ("change", ("Change ID",)),
    "Guidance and Handover": ("guide", ("Guide ID",)),
    "Source Bindings": ("binding", ("Binding ID",)), "Source Manifest": ("source", ("Source ID",)),
}
PROFILES = {
    "business": {"Requirements", "Stories", "Acceptance", "Roadmap", "Milestones", "Release Notes", "Test Runs", "Guidance and Handover"},
    "engineering": {"Requirements", "Stories", "Acceptance", "ADRs", "ADR Source Sections", "Design Views", "UX and Contracts", "Implementation Tasks", "Source Bindings", "Release Plan", "Release Notes", "Test Cases", "Test Runs"},
    "qa": {"Requirements", "Stories", "Acceptance", "Test Cases", "Test Runs", "Release Plan"},
    "operations": {"Release Plan", "Release Notes", "Guidance and Handover", "UX and Contracts", "Test Runs"},
    "executive": {"Roadmap", "Milestones", "Release Plan", "Release Notes"},
}
COMMON_TABLES = {"Artifact Register", "Sources", "Links", "Evidence", "Reviews and Approvals", "Risks and Decisions"}
PROVENANCE = ["Source Document", "Source Section", "Source Revision", "Source Hash"]


def _sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _json_bytes(data: object) -> bytes:
    return (json.dumps(data, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()


def _path(root: Path, value: Path | str, *, base: Path | None = None) -> Path:
    root = root.resolve()
    path = Path(value)
    path = (path if path.is_absolute() else (base or root) / path).resolve()
    if not path.is_relative_to(root):
        raise ArtifactError("Path escapes the declared root")
    if any(part == ".env" or part.startswith(".env.") or part == "secrets" for part in path.parts) or path.suffix in {".pem", ".key"}:
        raise ArtifactError("Credential/environment/key paths are outside this tool's scope")
    return path


def _read(root: Path, value: Path | str, *, base: Path | None = None) -> tuple[Path, bytes]:
    path = _path(root, value, base=base)
    if not path.is_file() or path.stat().st_size > MAX_SOURCE_BYTES:
        raise ArtifactError("Source is missing, not a file or too large")
    return path, path.read_bytes()


def _text(value: object) -> str:
    if not isinstance(value, str) or re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", value):
        raise ArtifactError("Source values must be literal strings without XML control characters")
    if len(value) > 32767:
        raise ArtifactError("Cell exceeds XLSX text limit; link the complete narrative instead")
    return value


class _UniqueLoader(yaml.SafeLoader):
    pass


def _mapping(loader, node):
    result = {}
    for key_node, value_node in node.value:
        key = loader.construct_object(key_node)
        if not isinstance(key, str) or key in result:
            raise ArtifactError("Duplicate or non-string metadata key")
        result[key] = loader.construct_object(value_node)
    return result


_UniqueLoader.add_constructor(yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, _mapping)


def parse_document(root: Path, source: Path, namespace: str) -> dict:
    path, raw = _read(root, source)
    content = raw.decode("utf-8")
    match = re.match(r"\A---\r?\n(.*?)\r?\n---(?:\r?\n|$)", content, re.S)
    if not match:
        raise ArtifactError("Instantiated source requires scalar YAML frontmatter")
    try:
        meta = yaml.load(match.group(1), Loader=_UniqueLoader)
    except yaml.YAMLError as exc:
        raise ArtifactError("Invalid or unsupported YAML frontmatter") from exc
    if not isinstance(meta, dict) or not META_FIELDS <= meta.keys():
        raise ArtifactError("Missing required document metadata")
    for value in meta.values():
        _text(value)
    if any(not meta[key].strip() for key in META_FIELDS):
        raise ArtifactError("Required metadata cannot be blank")
    if meta["contract_version"] != CONTRACT_VERSION:
        raise ArtifactError("Unsupported content contract version")
    if not namespace or meta["namespace"] != namespace:
        raise ArtifactError("Explicit namespace must match every document")
    if meta["status"] not in DOCUMENT_STATUSES or meta["document_type"] not in RECORD_TYPES:
        raise ArtifactError("Invalid document status or type")
    if not meta["source_authority"].startswith("authored Markdown"):
        raise ArtifactError("Declare authored Markdown authority; import engine fields explicitly")
    tables, narratives = {}, []
    body = content[match.end():]
    previous = ""
    for line in body.splitlines():
        if line.startswith("## ") and re.sub(r"^\d{2} ", "", line[3:]) in TABLE_HEADERS:
            if not re.fullmatch(r'<a id="[a-z0-9][a-z0-9-]*"></a>', previous):
                raise ArtifactError("Structured sections require explicit anchors")
        if line.strip():
            previous = line.strip()
    # Stable anchors, rather than headings or list positions, identify sections.
    for section in re.split(r'(?m)^<a id="([^"\n]+)"></a>\s*\n', body)[1:][::2]:
        if not re.fullmatch(r"[a-z0-9][a-z0-9-]*", section):
            raise ArtifactError("Invalid explicit section anchor")
    parts = re.split(r'(?m)^<a id="([^"\n]+)"></a>\s*\n', body)
    if parts[0].strip():
        narratives.append({"anchor": "", "text": _text(parts[0].strip())})
    anchors = set()
    for anchor, section in zip(parts[1::2], parts[2::2]):
        if anchor in anchors:
            raise ArtifactError("Duplicate section anchor")
        anchors.add(anchor)
        heading = re.search(r"^## (.+)$", section, re.M)
        title = re.sub(r"^\d{2} ", "", heading.group(1)) if heading else ""
        lines = [line for line in section.splitlines() if line.startswith("| ") and line.endswith(" |")]
        if title in TABLE_HEADERS:
            if title in tables or len(lines) < 3:
                raise ArtifactError("Duplicate or empty structured section")
            def cells(line):
                return [_text(v.strip().replace("<br>", "\n").replace("&#124;", "|"))
                        for v in line[1:-1].split("|")]
            headers = cells(lines[0])
            if headers != TABLE_HEADERS[title] or any(v != "---" for v in cells(lines[1])):
                raise ArtifactError("Structured section headers differ from contract")
            rows = []
            for line in lines[2:]:
                values = cells(line)
                if len(values) != len(headers):
                    raise ArtifactError("Malformed Markdown table row")
                if any(values):
                    rows.append(dict(zip(headers, values)))
            tables[title] = {"anchor": anchor, "rows": rows}
            narrative = "\n".join(line for line in section.splitlines() if line not in lines).strip()
            if narrative:
                narratives.append({"anchor": anchor, "text": _text(narrative)})
        else:
            narratives.append({"anchor": anchor, "text": _text(section.strip())})
    if re.search(r"^## (" + "|".join(re.escape(k) for k in TABLE_HEADERS) + r")\s*$", parts[0], re.M):
        raise ArtifactError("Structured sections require explicit anchors")
    return {"metadata": meta, "path": str(path.relative_to(root.resolve())),
            "sha256": _sha(raw), "tables": tables, "narratives": narratives}


@dataclass
class Bundle:
    root: Path
    namespace: str
    documents: list[dict]
    records: list[dict] = field(default_factory=list)
    sources: dict[str, dict] = field(default_factory=dict)
    issues: list[str] = field(default_factory=list)
    engine_limits: list[dict] = field(default_factory=list)


def _source(bundle: Bundle, path: Path, raw: bytes, authority: str, revision: str = "") -> None:
    relative = str(path.relative_to(bundle.root))
    if relative in bundle.sources and bundle.sources[relative]["sha256"] != _sha(raw):
        raise ArtifactError("Source changed during import")
    previous = bundle.sources.get(relative, {})
    if previous and previous["authority"] != authority:
        bundle.issues.append("Conflicting source authority declarations")
    if previous.get("revision") and revision and previous["revision"] != revision:
        bundle.issues.append("Conflicting source revision declarations")
    bundle.sources[relative] = {"path": relative, "sha256": _sha(raw), "authority": authority,
                               "revision": revision or previous.get("revision", "")}


def _engine(bundle: Bundle, path: Path) -> tuple[dict[str, dict], bytes]:
    path, raw = _read(bundle.root, path)
    try:
        data = json.loads(raw)
        items = data if isinstance(data, list) else [data]
        owners = {}
        for item in items:
            if not isinstance(item, dict):
                raise ArtifactError("Invalid engine source object")
            if "issue_id" in item:
                model, key, kind = IssueDraft.from_dict(item), "issue_id", "task"
            elif "story_id" in item:
                model, key, kind = UserStory.from_dict(item), "story_id", "story"
            elif "ready" in item:
                model, key, kind = ReadinessCheck.from_dict(item), "req_id", "readiness"
                if not isinstance(item["ready"], bool):
                    raise ArtifactError("Recorded readiness must have a boolean ready value")
            else:
                model, key, kind = RequirementPRD.from_dict(item), "req_id", "prd"
            native_id = _text(item[key])
            if not native_id or native_id in owners:
                raise ArtifactError("Duplicate or blank native engine ID")
            if getattr(model, "project", None) not in {None, bundle.namespace}:
                raise ArtifactError("Engine source project conflicts with explicit namespace")
            owners[native_id] = {"kind": kind, "data": item}
    except (KeyError, TypeError, ValueError) as exc:
        if isinstance(exc, ArtifactError):
            raise
        raise ArtifactError("Unsupported or invalid native engine schema") from exc
    _source(bundle, path, raw, "Engine JSON")
    return owners, raw


def _owner_values(owner: dict, source_field: str) -> list[str]:
    allowed = {"prd": {"functional_requirements", "non_functional_requirements", "standards_spec.acceptance_criteria"},
               "story": {"acceptance_criteria"}, "task": {"acceptance_criteria"}}
    if source_field not in allowed.get(owner["kind"], set()):
        raise ArtifactError("Unsupported binding field; raw engine metadata is never exported")
    values = owner["data"]
    for part in source_field.split("."):
        values = values[part]
    if not isinstance(values, list) or any(not isinstance(v, str) for v in values):
        raise ArtifactError("Binding field must be a string list")
    return [_text(value) for value in values]


def _record(document: dict, namespace: str, table: str, row: dict) -> dict:
    kind, keys = TABLE_IDENTITIES[table]
    if table == "Artifact Register":
        kind = row["Type"]
    if kind not in RECORD_TYPES or any(not row[key] for key in keys):
        raise ArtifactError("Record identity/type missing or invalid")
    rid = row[keys[0]] if len(keys) == 1 else json.dumps([row[k] for k in keys], separators=(",", ":"))
    return {"namespace": namespace, "record_type": kind, "record_id": rid,
            "table": table, "values": dict(row), "source_document": document["path"],
            "section": document["tables"][table]["anchor"],
            "revision": document["metadata"]["content_version"], "source_hash": document["sha256"]}


def load_bundle(root: Path, sources: list[Path], namespace: str, *, strict: bool = True) -> Bundle:
    root = root.resolve()
    if not re.fullmatch(r"[a-z][a-z0-9-]{0,62}", namespace or ""):
        raise ArtifactError("Explicit namespace must be a simple stable identifier")
    if not sources:
        raise ArtifactError("At least one source document is required")
    documents = [parse_document(root, path, namespace) for path in sources]
    bundle = Bundle(root, namespace, documents)
    seen, identities, engine_cache, native_owners = set(), set(), {}, {}
    for doc in documents:
        key = (namespace, doc["metadata"]["document_id"])
        if key in seen:
            raise ArtifactError("Duplicate source document identity")
        seen.add(key)
        identities.add((namespace, doc["metadata"]["document_type"], doc["metadata"]["document_id"]))
        path, raw = _read(root, doc["path"])
        if _sha(raw) != doc["sha256"]:
            raise ArtifactError("Source changed during import")
        _source(bundle, path, raw, "Authored Markdown", doc["metadata"]["content_version"])
        for table, spec in doc["tables"].items():
            for row in spec["rows"]:
                record = _record(doc, namespace, table, row)
                key = (record["table"], record["record_type"], record["record_id"])
                if key in seen:
                    raise ArtifactError("Duplicate record identity in a structured view")
                seen.add(key)
                identities.add((namespace, record["record_type"], record["record_id"]))
                if table == "Design Views":
                    identities.add((namespace, "design", row["Design ID"]))
                bundle.records.append(record)
    def get_engine(path):
        if path not in engine_cache:
            engine_cache[path] = _engine(bundle, path)
            for native_id, owner in engine_cache[path][0].items():
                key = (owner["kind"], native_id)
                if key in native_owners and native_owners[key] != owner["data"]:
                    raise ArtifactError("Conflicting native owner baselines")
                native_owners[key] = owner["data"]
        return engine_cache[path]
    bindings = {}
    for record in bundle.records:
        row = record["values"]
        base = (root / record["source_document"]).parent
        if record["table"] == "Source Manifest":
            path, raw = _read(root, row["Path"], base=base)
            if row["Source Kind"] == "Engine JSON":
                owners, raw = get_engine(path)
                for native_id, owner in owners.items():
                    if owner["kind"] == "readiness":
                        data = owner["data"]
                        bundle.engine_limits.append({"native_id": native_id,
                            "path": str(path.relative_to(root)),
                            "text": "Recorded engine readiness: status=" + data["status"] +
                                    "; ready=" + str(data["ready"]) +
                                    "; recorded source result only, not approval"})
            if row["SHA-256"] != _sha(raw):
                bundle.issues.append("Declared source manifest hash is stale")
            expected_authority = {"Engine JSON": "Engine JSON", "Authored Markdown": "Authored Markdown",
                                  "Machine source": "Machine source"}.get(row["Source Kind"])
            if expected_authority is None or row["Authority"] != expected_authority:
                bundle.issues.append("Unknown or conflicting source authority")
            _source(bundle, path, raw, row["Authority"], row["Revision"])
        if record["table"] == "Artifact Register" and row["Machine Source Link"]:
            path = _path(root, row["Machine Source Link"], base=base)
            if "Engine JSON" in row["Source Authority"]:
                owners, raw = get_engine(path)
                if row["Artifact ID"] not in owners:
                    bundle.issues.append("Registered native owner is missing")
                elif owners[row["Artifact ID"]]["kind"] != row["Type"]:
                    bundle.issues.append("Registered native type conflicts")
                else:
                    native = owners[row["Artifact ID"]]["data"]
                    record["values"]["Native Source Status"] = _text(native.get("status", ""))
                    record["values"]["Native Source Title"] = _text(native.get("title", ""))
                    record["values"]["Native PRD ID"] = _text(native.get("req_id", ""))
                    record["values"]["Native Story ID"] = _text(native.get("story_id", ""))
                _source(bundle, path, raw, "Engine JSON", row["Source Revision"])
            else:
                path, raw = _read(root, path)
                _source(bundle, path, raw, "Machine source", row["Source Revision"])
            if row["Source SHA-256"] != _sha(raw):
                bundle.issues.append("Registered machine source hash is stale")
        if record["table"] != "Source Bindings":
            continue
        if row["Binding Status"] == "retired":
            continue
        if row["Binding Status"] not in {"resolved", "proposed", "unresolved"}:
            raise ArtifactError("Invalid binding status")
        key = (row["Record Type"], row["Record ID"])
        if key in bindings:
            raise ArtifactError("Duplicate active source binding")
        if row["Source Kind"] != "Engine JSON":
            raise ArtifactError("Unsupported binding source kind")
        path = _path(root, row["Source Path"], base=base)
        owners, raw = get_engine(path)
        owner = owners.get(row["Native Owner ID"])
        if owner is None:
            raise ArtifactError("Binding native owner missing")
        values = _owner_values(owner, row["Source Field"])
        expected_type = "requirement" if "requirements" in row["Source Field"] else "criterion"
        if row["Record Type"] != expected_type:
            raise ArtifactError("Binding record type conflicts with source field")
        indices = [i for i, text in enumerate(values) if _sha(text.encode()) == row["Source Text SHA-256"]]
        same_source = _sha(raw) == row["Source File SHA-256"]
        locator = re.search(r"\[(\d+)\]$", row["Source Locator"])
        position = int(locator.group(1)) if locator else None
        # Resolved duplicate text can use the explicitly reviewed original locator
        # only at the exact reviewed file hash. Drift makes it ambiguous again.
        index = indices[0] if len(indices) == 1 else (
            position if same_source and row["Binding Status"] == "resolved" and position in indices else None)
        if not same_source or row["Binding Status"] != "resolved" or index is None:
            bundle.issues.append("Source binding is unresolved or stale")
        bindings[key] = {"text": values[index] if index is not None else None,
                         "owner": owner, "row": row}
    for record in bundle.records:
        key = (record["record_type"], record["record_id"])
        binding = bindings.get(key)
        if binding:
            value = binding["text"]
            record["values"]["Native Source Text"] = value or "unresolved source association"
            record["values"]["Native Owner ID"] = binding["row"]["Native Owner ID"]
            if record["table"] == "Requirements" and value is not None:
                if record["values"]["Statement"] != value:
                    bundle.issues.append("Authored requirement differs from engine-owned statement")
                record["values"]["Statement"] = value
        if record["table"] == "Stories":
            for owners, _ in engine_cache.values():
                owner = owners.get(record["record_id"])
                if owner and owner["kind"] == "story":
                    for column, field_name in {"Actor": "actor", "Goal": "want", "Benefit": "outcome"}.items():
                        native = _text(owner["data"][field_name])
                        if record["values"][column] != native:
                            bundle.issues.append("Authored story differs from engine-owned field")
                        record["values"][column] = native
    for kind, rid in bindings:
        if (namespace, kind, rid) not in identities:
            bundle.issues.append("Binding references an undeclared record identity")
    _validate_records(bundle, identities)
    if strict and bundle.issues:
        raise ArtifactError("; ".join(sorted(set(bundle.issues))))
    return bundle


def _validate_records(bundle: Bundle, identities: set) -> None:
    evidence = {r["record_id"]: r["values"] for r in bundle.records if r["table"] == "Evidence"}
    def observed(eid):
        item = evidence.get(eid, {})
        return all(item.get(name) and item[name] != "unconfirmed"
                   for name in ("URI / Path", "Revision", "Observed Result", "Observed UTC")) and timestamp(item["Observed UTC"])
    def timestamp(value):
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00")).tzinfo is not None
        except ValueError:
            return False
    for record in bundle.records:
        row, table = record["values"], record["table"]
        if row.get("Source Artifact ID") and not any(key[0] == bundle.namespace and key[2] == row["Source Artifact ID"] for key in identities):
            bundle.issues.append("Source artifact reference unresolved")
        if table == "Requirements" and row["Source ID"] and (bundle.namespace, "source", row["Source ID"]) not in identities:
            bundle.issues.append("Requirement source reference unresolved")
        if table == "Links":
            for end in ("From", "To"):
                key = tuple(row[end + " " + suffix] for suffix in ("Namespace", "Type", "ID"))
                if key not in identities:
                    bundle.issues.append("Link endpoint unresolved; supply its declared source baseline")
            if row["Relation"] not in {"derived_from", "satisfies", "verifies", "included_in", "supersedes", "depends_on", "refers_to"}:
                bundle.issues.append("Invalid link relation")
        if table == "Requirements" and row["Class"] not in {"Business", "Stakeholder", "Solution", "Transition"}:
            bundle.issues.append("Invalid requirement class")
        if table == "Requirements":
            if not row["Statement"]:
                bundle.issues.append("Requirement statement missing")
            if row["Agreed Priority"] not in {"", "unconfirmed"} and not observed(row["Decision Evidence"]):
                bundle.issues.append("Agreed priority lacks decision evidence")
        if table == "Acceptance" and any(not row[name] for name in ("Context", "Event", "Expected Outcome", "Verification Method")):
            bundle.issues.append("Acceptance criterion lacks observable condition/method")
        if table == "ADRs":
            if row["Status"] not in {"Proposed", "Accepted", "Superseded"}:
                bundle.issues.append("Invalid ADR decision status")
            if row["Status"] == "Accepted" and (row["Decision Owner"] in {"", "unconfirmed"} or not observed(row["Decision Evidence"])):
                bundle.issues.append("Accepted ADR lacks decision evidence")
        if table == "Test Cases" and (bundle.namespace, "criterion", row["Criterion ID"]) not in identities:
            bundle.issues.append("Test criterion reference unresolved")
        if table == "Release Notes" and not observed(row["Evidence ID"]):
            bundle.issues.append("Delivered release note lacks observed evidence")
        if table == "Test Runs":
            if (bundle.namespace, "test_case", row["Test ID"]) not in identities:
                bundle.issues.append("Test case reference unresolved")
            if row["Result"] not in {"Not run", "Blocked", "Skipped", "Fail", "Pass"}:
                bundle.issues.append("Invalid test result")
            if row["Result"] in {"Pass", "Fail"}:
                fields = ("Actual Observation", "Revision", "Environment", "Executed UTC", "Executor", "Evidence ID")
                if any(not row[field] or row[field] == "unconfirmed" for field in fields) or not observed(row["Evidence ID"]) or not timestamp(row["Executed UTC"]):
                    bundle.issues.append("Executed test result lacks scoped evidence")
        if table == "Reviews and Approvals" and row["Review Decision"] not in {"Not reviewed", "Agree with proposal", "Requires revision", "Needs decision"}:
            bundle.issues.append("Invalid reviewer feedback value")
    for doc in bundle.documents:
        if doc["metadata"]["status"] != "accepted":
            continue
        reviews = doc["tables"].get("Reviews and Approvals", {}).get("rows", [])
        if doc["metadata"]["owner"] == "unconfirmed" or not any(
            row["Source Artifact ID"] == doc["metadata"]["document_id"] and
            row["Source Revision"] == doc["metadata"]["content_version"] and
            row["Review Decision"] == "Agree with proposal" and
            row["Reviewer / Role"] not in {"", "unconfirmed"} and timestamp(row["Reviewed UTC"]) and
            observed(row["Evidence ID"]) for row in reviews
        ):
            bundle.issues.append("Accepted document lacks scoped reviewer evidence")


def propose_bindings(root: Path, source: Path, namespace: str, *, registry: Path | None = None,
                     output: Path | None = None) -> str:
    """Return a new Markdown proposal. Never marks new or changed bindings resolved."""
    root = root.resolve()
    if not re.fullmatch(r"[a-z][a-z0-9-]{0,62}", namespace or ""):
        raise ArtifactError("Explicit namespace required")
    bundle = Bundle(root, namespace, [])
    source = _path(root, source)
    owners, raw = _engine(bundle, source)
    prior, retired = [], []
    if registry is not None:
        doc = parse_document(root, registry, namespace)
        for row in doc["tables"].get("Source Bindings", {}).get("rows", []):
            old_path = _path(root, row["Source Path"], base=(root / doc["path"]).parent)
            if old_path == source:
                (retired if row["Binding Status"] == "retired" else prior).append(row)
    output = _path(root, output) if output else source.parent / "binding-proposal.md"
    rows, used = [], set()
    for native_id, owner in owners.items():
        fields = ["functional_requirements", "non_functional_requirements", "standards_spec.acceptance_criteria"] if owner["kind"] == "prd" else ["acceptance_criteria"]
        if owner["kind"] == "readiness":
            continue
        for field_name in fields:
            values = _owner_values(owner, field_name)
            current_hashes = {_sha(value.encode()) for value in values}
            for index, value in enumerate(values):
                text_hash = _sha(value.encode())
                candidates = [r for r in prior if r["Native Owner ID"] == native_id and r["Source Field"] == field_name and r["Source Text SHA-256"] == text_hash]
                old = candidates[0] if len(candidates) == 1 else None
                if old is None:
                    candidates = [r for r in prior if r["Native Owner ID"] == native_id and
                                  r["Source Field"] == field_name and
                                  (r["Source Text SHA-256"] == text_hash or r["Source Text SHA-256"] not in current_hashes) and
                                  re.search(rf"\[{index}\]$", r["Source Locator"])]
                    old = candidates[0] if len(candidates) == 1 else None
                kind = "requirement" if "requirements" in field_name else "criterion"
                rid = old["Record ID"] if old else ("REQ-" if kind == "requirement" else "AC-") + uuid.uuid4().hex[:12]
                if rid in used:
                    if values.count(value) > 1:
                        old = None
                        rid = ("REQ-" if kind == "requirement" else "AC-") + uuid.uuid4().hex[:12]
                    else:
                        raise ArtifactError("Ambiguous prior registry identities; review mappings manually")
                used.add(rid)
                ambiguous = values.count(value) > 1 or (old and old["Source Text SHA-256"] != text_hash)
                state = "unresolved" if ambiguous else "proposed"
                rows.append([old["Binding ID"] if old else "BIND-" + uuid.uuid4().hex[:12], kind, rid, native_id,
                             os.path.relpath(source, output.parent), "Engine JSON", field_name, f"[{index}]",
                             text_hash, _sha(raw), state])
    # Missing entries remain visible proposals, not silent removals or automatic
    # retirements; preserve existing retired history as well.
    columns = TABLE_HEADERS["Source Bindings"]
    for old in [row for row in prior if row["Record ID"] not in used] + retired:
        preserved = dict(old)
        preserved["Source Path"] = os.path.relpath(source, output.parent)
        if preserved["Binding Status"] != "retired":
            preserved["Binding Status"] = "unresolved"
        rows.append([preserved[name] for name in columns])
    meta = {"contract_version": CONTRACT_VERSION, "namespace": namespace,
            "document_id": "DOC-" + uuid.uuid4().hex[:12], "document_type": "document",
            "content_version": "0.1", "status": "draft", "owner": "unconfirmed",
            "requested_action": "review binding proposal", "source_authority": "authored Markdown"}
    header = ["---", *[key + ": " + json.dumps(value) for key, value in meta.items()], "---", "",
              "# Binding proposal", "", "Review associations before persisting resolved mappings. This file changes no source.", "",
              '<a id="source-bindings"></a>', "## Source Bindings", ""]
    columns = TABLE_HEADERS["Source Bindings"]
    header += ["| " + " | ".join(columns) + " |", "| " + " | ".join(["---"] * len(columns)) + " |"]
    header += ["| " + " | ".join(v.replace("|", "&#124;").replace("\n", "<br>") for v in row) + " |" for row in rows]
    if not rows:
        header.append("| " + " | ".join([""] * len(columns)) + " |")
    return "\n".join(header) + "\n"


def _spreadsheet():
    try:
        import openpyxl
    except ImportError as exc:
        raise ArtifactError("Install the optional workbooks extra for XLSX operations") from exc
    return openpyxl


def _project(bundle: Bundle, audience: str, baseline_revision: str) -> dict:
    if audience not in PROFILES or not baseline_revision.strip():
        raise ArtifactError("Known audience and explicit baseline revision required")
    selected = PROFILES[audience] | COMMON_TABLES
    grouped = {}
    for record in bundle.records:
        if record["table"] in selected:
            grouped.setdefault(record["table"], []).append(record)
    tables = {}
    for title, records in grouped.items():
        fields = list(TABLE_HEADERS[title])
        for record in records:
            for name in record["values"]:
                if name not in fields:
                    fields.append(name)
        headers = ["Namespace", "Record Type", "Record ID", *fields, *PROVENANCE, "Feedback"]
        rows = [[record["namespace"], record["record_type"], record["record_id"],
                 *[record["values"].get(name, "") for name in fields],
                 record["source_document"], record["section"], record["revision"], record["source_hash"], ""]
                for record in records]
        tables[title] = {"headers": headers, "rows": rows}
    # Full Markdown narrative is available to engineering; other profiles retain
    # document identity and full-source links, rather than silently truncating it.
    if audience == "engineering":
        rows = []
        for doc in bundle.documents:
            for section in doc["narratives"]:
                rows.append([bundle.namespace, "document_section",
                             doc["metadata"]["document_id"] + "#" + section["anchor"],
                             section["text"], doc["path"], section["anchor"],
                             doc["metadata"]["content_version"], doc["sha256"], ""])
        if rows:
            tables["Narrative"] = {"headers": ["Namespace", "Record Type", "Record ID", "Markdown Content", *PROVENANCE, "Feedback"], "rows": rows}
    present = {record["table"] for record in bundle.records}
    omissions = sorted(present - selected)
    if audience != "engineering" and any(doc["narratives"] for doc in bundle.documents):
        omissions.append("Full narrative: use complete Markdown source links")
    export_id = "EXPORT-" + uuid.uuid4().hex
    control = [("export_id", export_id), ("contract_version", CONTRACT_VERSION), ("profile_version", "0.1"),
               ("audience", audience), ("baseline_revision", baseline_revision),
               ("generated_utc", datetime.now(timezone.utc).isoformat()),
               ("omitted_detail", "; ".join(omissions) or "none"),
               ("source_validation", "structural integrity only; no review, execution or approval granted")]
    for doc in bundle.documents:
        control.append(("document:" + doc["metadata"]["document_id"],
                        doc["path"] + " | " + doc["metadata"]["status"] + " | " + doc["metadata"]["requested_action"]))
    tables["Control"] = {"headers": ["Namespace", "Record Type", "Record ID", "Value", "Feedback"],
                         "rows": [[bundle.namespace, "control", key, value, ""] for key, value in control]}
    tables["Source Manifest"] = {"headers": ["Namespace", "Record Type", "Record ID", "Source Document", "Authority", "SHA-256", "Declared Source Revision", "Baseline Revision", "Feedback"],
        "rows": [[bundle.namespace, "source_file", source["path"], source["path"], source["authority"], source["sha256"], source["revision"] or "not recorded", baseline_revision, ""]
                 for source in sorted(bundle.sources.values(), key=lambda value: value["path"])]}
    # Readiness/blocker facts survive every profile, even if their detailed view
    # is omitted. This does not turn missing information into a false pass.
    limits = [[bundle.namespace, "limit", doc["metadata"]["document_id"],
               "Document status: " + doc["metadata"]["status"] + "; source/export validation is not approval.",
               doc["path"], ""] for doc in bundle.documents]
    limits.extend([bundle.namespace, "recorded_readiness", item["native_id"], item["text"], item["path"], ""]
                  for item in bundle.engine_limits)
    for record in bundle.records:
        if record["table"] in {"Risks and Decisions", "Test Runs", "Release Plan"}:
            limits.append([bundle.namespace, record["record_type"], record["table"] + ":" + record["record_id"],
                           " | ".join(key + ": " + value for key, value in record["values"].items() if value),
                           record["source_document"], ""])
    tables["Readiness and Blockers"] = {"headers": ["Namespace", "Record Type", "Record ID", "Known State / Limit", "Source Document", "Feedback"], "rows": limits}
    tables["Read Me"] = {"headers": ["Namespace", "Record Type", "Record ID", "Guidance", "Feedback"],
        "rows": [[bundle.namespace, "guidance", "authority", "Markdown and declared engine sources are canonical. This is a derived audience view.", ""],
                 [bundle.namespace, "guidance", "feedback", "Edit/comment for review; compare the returned workbook against the original XLSX and JSON baseline. Feedback never applies source changes or approvals.", ""],
                 [bundle.namespace, "guidance", "freshness", "Check every source-manifest hash before relying on currentness. Preserve annotated workbooks and generate a new export.", ""]]}
    order = ["Read Me", "Control", "Source Manifest", "Readiness and Blockers", *grouped]
    if "Narrative" in tables:
        order.append("Narrative")
    return {"contract_version": CONTRACT_VERSION, "export_id": export_id, "namespace": bundle.namespace,
            "audience": audience, "profile_version": "0.1", "baseline_revision": baseline_revision,
            "source_documents": [doc["path"] for doc in bundle.documents],
            "sources": list(bundle.sources.values()), "omissions": omissions,
            "tables": {title: tables[title] for title in dict.fromkeys(order)}}


def write_new(root: Path, outputs: dict[Path, bytes]) -> None:
    """Reserve distinct new outputs before writing. Never replace existing paths."""
    paths = [_path(root, path) for path in outputs]
    if len(set(paths)) != len(paths):
        raise ArtifactError("Output paths must be distinct")
    if any(path.exists() for path in paths):
        raise FileExistsError("Refusing to overwrite an existing output")
    opened = []
    try:
        for path in paths:
            path.parent.mkdir(parents=True, exist_ok=True)
            opened.append((path, path.open("xb")))
        for (_, stream), data in zip(opened, outputs.values()):
            stream.write(data)
            stream.flush()
    except Exception:
        # Only remove our still-empty reservations; never remove an annotated or
        # replaced file, or claim a partially written bundle was completed.
        for path, stream in opened:
            if path.exists() and path.stat().st_ino == os.fstat(stream.fileno()).st_ino and path.stat().st_size == 0:
                path.unlink()
        raise
    finally:
        for _, stream in opened:
            stream.close()


def export_workbook(bundle: Bundle, audience: str, output: Path, snapshot: Path,
                    baseline_revision: str) -> dict:
    if _path(bundle.root, output) == _path(bundle.root, snapshot):
        raise ArtifactError("Workbook and snapshot paths must be distinct")
    if bundle.issues:
        raise ArtifactError("Resolve source integrity issues before export")
    for source in bundle.sources.values():
        _, raw = _read(bundle.root, source["path"])
        if source["sha256"] != _sha(raw):
            raise ArtifactError("Source changed since validation")
    payload = _project(bundle, audience, baseline_revision)
    scan_cells = sum((4 + max(1, len(table["rows"]))) * len(table["headers"])
                     for table in payload["tables"].values())
    if scan_cells > MAX_WORKBOOK_SCAN_CELLS:
        raise ArtifactError("Workbook exceeds scan cell limit; split the source set")
    digest = _sha(_json_bytes(payload))
    openpyxl = _spreadsheet()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Engineering artifacts: " + audience + " view"
    workbook.properties.creator = "Engineering documentation"
    workbook.properties.identifier = "sha256:" + digest
    workbook.properties.description = "Derived view; proposed feedback only; no approval or execution implied"
    output = _path(bundle.root, output)
    for number, (title, table) in enumerate(payload["tables"].items()):
        ws = workbook.create_sheet(title)
        ws["A1"] = title
        ws["A1"].font = Font(name="Arial", bold=True, size=12)
        ws["D1"] = "Derived " + audience + " view"
        ws["A2"], ws["D2"] = "Baseline", baseline_revision
        ws["A3"], ws["D3"] = "Navigation", "Return to Read Me"
        ws["D3"].hyperlink = "#'Read Me'!A1"
        headers = table["headers"]
        rows = table["rows"] or [[""] * len(headers)]
        for row_number, row in enumerate([headers, *rows], 4):
            lines = 1
            for column, value in enumerate(row, 1):
                value = _text(value)
                cell = ws.cell(row_number, column, value)
                cell.data_type = "s"
                cell.font = Font(name="Arial", size=11, bold=row_number == 4,
                                 color="FFFFFF" if row_number == 4 else "172B3A")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if row_number == 4:
                    cell.fill = PatternFill("solid", fgColor="16324F")
                elif headers[column - 1] == "Feedback":
                    cell.fill = PatternFill("solid", fgColor="FFF2CC")
                elif headers[column - 1] == "Source Document" and value:
                    source_path = _path(bundle.root, value)
                    cell.hyperlink = os.path.relpath(source_path, output.parent)
                lines = max(lines, sum(max(1, (len(part) + 49) // 50) for part in value.splitlines()))
            ws.row_dimensions[row_number].height = max(34, min(400, 16 * lines + 8))
        for index, header in enumerate(headers, 1):
            width = 18 if index <= 3 else 65 if "Hash" in header or "SHA-256" in header else 48
            ws.column_dimensions[get_column_letter(index)].width = width
        table_object = Table(displayName=f"EngineeringView{number:02d}", ref=f"A4:{get_column_letter(len(headers))}{4 + len(rows)}")
        table_object.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table_object)
        ws.freeze_panes = "D5"
        ws.sheet_view.zoomScale = 85
        ws.print_title_rows = "1:4"
        ws.print_title_cols = "A:C"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 2, 0
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook_bytes = buffer.getvalue()
    ledger = {"payload": payload, "payload_sha256": digest, "workbook_sha256": _sha(workbook_bytes)}
    ledger_bytes = _json_bytes(ledger)
    if len(workbook_bytes) > MAX_WORKBOOK_BYTES or len(ledger_bytes) > MAX_SOURCE_BYTES:
        raise ArtifactError("Export exceeds supported workbook/snapshot size; split the source set")
    write_new(bundle.root, {output: workbook_bytes, snapshot: ledger_bytes})
    return ledger


def _read_workbook(root: Path, source: Path):
    path = _path(root, source)
    if not path.is_file() or path.stat().st_size > MAX_WORKBOOK_BYTES:
        raise ArtifactError("Workbook missing or exceeds supported size")
    raw = path.read_bytes()
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 64_000_000:
                raise ArtifactError("Workbook expanded size exceeds supported limit")
            if any("vbaProject" in item.filename or item.filename.startswith("xl/externalLinks/") for item in entries):
                raise ArtifactError("Macros and external calculation links are unsupported; original is preserved")
        workbook = _spreadsheet().load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ArtifactError("Unsupported workbook transport") from exc
    # Sparse cells can imply a huge rectangular scan despite a tiny ZIP payload.
    if sum(ws.max_row * ws.max_column for ws in workbook.worksheets) > MAX_WORKBOOK_SCAN_CELLS:
        raise ArtifactError("Workbook exceeds scan cell limit; split or trim the review range")
    return workbook, raw


def _sheet_records(ws, headers: list[str]) -> tuple[dict, list]:
    problems, rows = [], {}
    actual = [ws.cell(4, index + 1).value for index in range(max(ws.max_column, len(headers)))]
    if actual != headers:
        problems.append({"kind": "headers_changed", "sheet": ws.title})
        return rows, problems
    for cells in ws.iter_rows(min_row=5, max_col=len(headers)):
        values = ["" if cell.value is None else str(cell.value) for cell in cells]
        key = tuple(values[:3])
        for cell in cells:
            if cell.comment:
                problems.append({"kind": "cell_comment", "sheet": ws.title, "identity": list(key),
                                 "cell": cell.coordinate, "author": cell.comment.author, "text": cell.comment.text})
        if not any(values):
            continue
        if any(not value for value in key) or key in rows:
            problems.append({"kind": "duplicate_or_blank_identity", "sheet": ws.title, "identity": list(key)})
            continue
        rows[key] = values
        for cell in cells:
            if cell.data_type == "f":
                problems.append({"kind": "formula_input", "sheet": ws.title, "cell": cell.coordinate})
    return rows, problems


def feedback_report(root: Path, baseline: Path, snapshot: Path, returned: list[Path]) -> dict:
    root = root.resolve()
    _, raw = _read(root, snapshot)
    ledger = json.loads(raw)
    payload = ledger["payload"]
    digest = _sha(_json_bytes(payload))
    original, original_raw = _read_workbook(root, baseline)
    if digest != ledger["payload_sha256"] or _sha(original_raw) != ledger["workbook_sha256"] or original.properties.identifier != "sha256:" + digest:
        raise ArtifactError("Baseline workbook/snapshot integrity mismatch")
    for title, table in payload["tables"].items():
        if title not in original.sheetnames:
            raise ArtifactError("Original baseline worksheet missing")
        rows, problems = _sheet_records(original[title], table["headers"])
        expected = {tuple(row[:3]): row for row in table["rows"]}
        if problems or rows != expected:
            raise ArtifactError("Original baseline cells differ from snapshot")
    stale = []
    for source in payload["sources"]:
        try:
            _, current_bytes = _read(root, source["path"])
            if _sha(current_bytes) != source["sha256"]:
                stale.append(source["path"])
        except ArtifactError:
            stale.append(source["path"])
    current_tables, issues = {}, []
    try:
        current = load_bundle(root, [Path(path) for path in payload["source_documents"]], payload["namespace"], strict=False)
        current_tables = _project(current, payload["audience"], payload["baseline_revision"])["tables"]
        issues = sorted(set(current.issues))
    except (ArtifactError, OSError, UnicodeError, ValueError) as exc:
        issues = ["Current source cannot be validated: " + type(exc).__name__]
    changes, proposals = [], {}
    for returned_path in returned:
        workbook, _ = _read_workbook(root, returned_path)
        returned_name = str(_path(root, returned_path).relative_to(root))
        for title in set(workbook.sheetnames) - set(payload["tables"]):
            changes.append({"kind": "worksheet_added", "sheet": title, "workbook": returned_name})
        for title, table in payload["tables"].items():
            if title not in workbook.sheetnames:
                changes.append({"kind": "worksheet_removed", "sheet": title, "workbook": returned_name})
                continue
            received, problems = _sheet_records(workbook[title], table["headers"])
            changes.extend({**problem, "workbook": returned_name} for problem in problems)
            old = {tuple(row[:3]): row for row in table["rows"]}
            current_table = current_tables.get(title, {})
            live = {tuple(row[:3]): dict(zip(current_table["headers"], row)) for row in current_table.get("rows", [])}
            for key in old.keys() - received.keys():
                baseline_fields = dict(zip(table["headers"], old[key]))
                changed = key in live and any(live[key].get(name) != value for name, value in baseline_fields.items())
                changes.append({"kind": "row_removed", "sheet": title, "identity": list(key), "workbook": returned_name,
                                "status": "conflict" if changed or issues else "proposed_stale" if stale else "proposed"})
            for key in received.keys() - old.keys():
                added_fields = dict(zip(table["headers"], received[key]))
                collision = key in live and any(live[key].get(name) != value for name, value in added_fields.items())
                changes.append({"kind": "row_added", "sheet": title, "identity": list(key), "values": received[key], "workbook": returned_name,
                                "status": "conflict" if collision or issues else "proposed_stale" if stale else "proposed"})
            for key in old.keys() & received.keys():
                for column, name in enumerate(table["headers"]):
                    previous, proposed = old[key][column], received[key][column]
                    if previous == proposed:
                        continue
                    current_value = live.get(key, {}).get(name)
                    status = "conflict" if (current_value is None or
                        (current_value != previous and current_value != proposed) or
                        (issues and name != "Feedback")) else "proposed_stale" if stale else "proposed"
                    change = {"kind": "field_change", "sheet": title, "identity": list(key), "field": name,
                              "baseline": previous, "current": current_value, "proposed": proposed,
                              "status": status, "workbook": returned_name}
                    changes.append(change)
                    proposals.setdefault((title, key, name), []).append(change)
    for related in proposals.values():
        if len({change["proposed"] for change in related}) > 1:
            for change in related:
                change["status"] = "conflict"
                change["reason"] = "Conflicting returned workbooks"
    row_proposals = {}
    for change in changes:
        if change.get("kind") in {"row_added", "row_removed", "field_change"}:
            row_proposals.setdefault((change["sheet"], tuple(change["identity"])), []).append(change)
    for related in row_proposals.values():
        kinds = {change["kind"] for change in related}
        added = {json.dumps(change["values"]) for change in related if change["kind"] == "row_added"}
        if ("row_removed" in kinds and kinds != {"row_removed"}) or len(added) > 1:
            for change in related:
                change["status"] = "conflict"
                change["reason"] = "Conflicting row removal/addition or modification"
    return {"export_id": payload["export_id"], "scope": "Proposed feedback only; no source, approval or workbook changes applied",
            "stale_sources": stale, "current_source_issues": issues, "changes": changes}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path.cwd(), help="Declared local source/output root")
    commands = parser.add_subparsers(dest="command", required=True)
    for name in ("validate", "export"):
        command = commands.add_parser(name)
        command.add_argument("--source", action="append", type=Path, required=True)
        command.add_argument("--namespace", required=True)
        if name == "export":
            command.add_argument("--audience", choices=sorted(PROFILES), required=True)
            command.add_argument("--output", type=Path, required=True)
            command.add_argument("--snapshot", type=Path, required=True)
            command.add_argument("--baseline-revision", required=True)
    bindings = commands.add_parser("bindings")
    bindings.add_argument("--source", type=Path, required=True)
    bindings.add_argument("--namespace", required=True)
    bindings.add_argument("--registry", type=Path)
    bindings.add_argument("--output", type=Path, required=True)
    feedback = commands.add_parser("feedback")
    feedback.add_argument("--baseline", type=Path, required=True)
    feedback.add_argument("--snapshot", type=Path, required=True)
    feedback.add_argument("--returned", type=Path, action="append", required=True)
    feedback.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        if args.command in {"validate", "export"}:
            bundle = load_bundle(args.root, args.source, args.namespace)
            if args.command == "export":
                export_workbook(bundle, args.audience, args.output, args.snapshot, args.baseline_revision)
                print("Created new audience workbook and baseline snapshot; no source changes")
            else:
                print(f"Source integrity valid: {len(bundle.records)} records; structural check only")
        elif args.command == "bindings":
            content = propose_bindings(args.root, args.source, args.namespace, registry=args.registry, output=args.output)
            write_new(args.root, {args.output: content.encode()})
            print("Created new binding proposal; review required; no source changes")
        else:
            report = feedback_report(args.root, args.baseline, args.snapshot, args.returned)
            write_new(args.root, {args.output: _json_bytes(report)})
            print(f"Created feedback report: {len(report['changes'])} observations; no changes applied")
        return 0
    except (ArtifactError, FileExistsError, OSError, UnicodeError, ValueError, KeyError) as exc:
        print("Operation refused: " + (str(exc) if isinstance(exc, (ArtifactError, FileExistsError)) else type(exc).__name__))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
